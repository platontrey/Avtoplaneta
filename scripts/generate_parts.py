import sys
import os
from openpyxl import load_workbook

def get_realistic_price(name, category):
    name_lower = name.lower()
    category_lower = category.lower()

    # Болты
    if 'болт' in name_lower:
        return 100

    # Двигатель и крупные компоненты
    if 'двигатель' in name_lower or 'блок цилиндров' in name_lower:
        return 100000

    # Трансмиссия
    if category_lower == 'трансмиссия':
        if 'акпп' in name_lower or 'мкпп' in name_lower:
            return 50000
        if 'дифференциал' in name_lower:
            return 20000
        return 5000

    # Подвеска
    if 'подвеска' in category_lower or 'подвеска' in name_lower:
        if 'амортизатор' in name_lower:
            return 2000
        if 'рычаг' in name_lower or 'кулак' in name_lower:
            return 3000
        return 1000

    # Тормозная система
    if 'тормоз' in category_lower or 'тормоз' in name_lower:
        if 'диск' in name_lower:
            return 1500
        if 'суппорт' in name_lower:
            return 2500
        return 500

    # Электрооснащение
    if category_lower == 'электрооснащение':
        if 'аккумулятор' in name_lower:
            return 5000
        if 'генератор' in name_lower or 'стартер' in name_lower:
            return 8000
        if 'блок управления' in name_lower:
            return 3000
        return 1000

    # Оптика
    if category_lower == 'оптика':
        if 'фара' in name_lower:
            return 3000
        return 500

    # Кузов
    if 'кузов' in category_lower:
        if 'дверь' in name_lower or 'крыло' in name_lower or 'капот' in name_lower:
            return 5000
        if 'стекло' in name_lower:
            return 2000
        return 1000

    # Система охлаждения
    if 'охлаждени' in category_lower:
        if 'радиатор' in name_lower:
            return 3000
        if 'помпа' in name_lower:
            return 2000
        return 500

    # Диски и шины
    if 'диски и шины' in category_lower:
        if 'шина' in name_lower:
            return 3000
        if 'диск' in name_lower:
            return 5000
        return 200

    # По умолчанию
    return 1000

def parse_xlsx_to_parts(xlsx_file):
    wb = load_workbook(xlsx_file)
    ws = wb.active  # Первый лист

    parts = []

    for row in ws.iter_rows(min_row=2, values_only=True):  # Пропустить заголовок
        if not row or not row[1] or not row[4]:  # Пропустить пустые строки или без имени/категории
            continue

        # Столбцы: B=1, E=4, I=8, J=9, K=10, V=21
        name = row[1] if row[1] else ''
        category = row[4] if row[4] else ''
        description = row[21] if len(row) > 21 and row[21] else name  # Использовать V для описания, fallback на имя
        front_rear = row[8] if len(row) > 8 and row[8] else ''
        left_right = row[9] if len(row) > 9 and row[9] else ''
        top_bottom_raw = row[10] if len(row) > 10 and row[10] else ''
        if top_bottom_raw.lower() == 'верхняя':
            top_bottom = 'Верх'
        elif top_bottom_raw.lower() == 'нижняя':
            top_bottom = 'Низ'
        else:
            top_bottom = top_bottom_raw.capitalize()

        price = get_realistic_price(name, category)

        parts.append({
            'name': name,
            'category': category,
            'description': description,
            'front_rear': front_rear,
            'left_right': left_right,
            'top_bottom': top_bottom,
            'quantity': 0,
            'price': price
        })

    return parts

def generate_js_code(parts):
    lines = []
    for part in parts:
        line = f"       {{ name: \"{part['name']}\", category: \"{part['category']}\", description: \"{part['description']}\", quantity: {part['quantity']}, price: {part['price']}"
        if part['front_rear']:
            line += f", front_rear: \"{part['front_rear']}\""
        if part['left_right']:
            line += f", left_right: \"{part['left_right']}\""
        if part['top_bottom']:
            line += f", top_bottom: \"{part['top_bottom']}\""
        line += " },"
        lines.append(line)
    return "\n".join(lines)

def main():
    if len(sys.argv) < 2 or len(sys.argv) > 3:
        print("Использование: python generate_parts.py <xlsx_file> [output_file]")
        sys.exit(1)

    xlsx_file = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) == 3 else None

    if not os.path.exists(xlsx_file):
        print(f"Файл {xlsx_file} не найден")
        sys.exit(1)

    parts = parse_xlsx_to_parts(xlsx_file)
    js_code = generate_js_code(parts)

    result = "// Сгенерированные данные запчастей\n" + js_code

    if output_file:
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write(result)
        print(f"Результат сохранён в {output_file}")
    else:
        print(result)

if __name__ == "__main__":
    main()